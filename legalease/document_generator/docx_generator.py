"""
LegalEase DOCX Generator
Produces beautifully styled, professional Microsoft Word (.docx) legal documents
with custom branding, typography, tables, and signature blocks.
"""
import os
import io
import re
from typing import Optional, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def _set_cell_background(cell, fill_hex: str):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cell (in twips: 20 twips = 1 pt)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def _set_cell_borders(cell, border_color="CCCCCC"):
    """Set subtle border around table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)


def generate_docx(
    title: str,
    content: str,
    summary_table: Optional[Dict[str, str]] = None,
    output_path: Optional[str] = None,
    branding: Optional[Dict[str, Any]] = None
) -> bytes:
    """
    Generate a styled, corporate-grade .docx legal contract.
    """
    doc = Document()
    branding = branding or {}

    font_name = branding.get("font_family", "Calibri")
    company_name = branding.get("company_name", "LegalEase AI Document Drafting")
    logo_path = branding.get("logo_path")
    header_text = branding.get("header_text", company_name)
    footer_text = branding.get("footer_text", "Confidential Legal Document — Generated with LegalEase")
    include_signatures = branding.get("include_signatures", True)

    # Palette
    NAVY_HEX = "0B2545"
    SLATE_HEX = "475569"
    LIGHT_BG_HEX = "F1F5F9"
    BORDER_HEX = "CBD5E1"
    
    PRIMARY_COLOR = RGBColor(11, 37, 69)      # Deep Navy
    SECONDARY_COLOR = RGBColor(71, 85, 105)  # Slate
    BODY_COLOR = RGBColor(30, 41, 59)        # Dark Charcoal

    # Document Page Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(header_text.upper())
        hrun.font.name = font_name
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = SECONDARY_COLOR

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(footer_text)
        frun.font.name = font_name
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = SECONDARY_COLOR

    # Add Optional Brand Logo / Top Header Block
    if logo_path and os.path.exists(logo_path):
        try:
            doc.add_picture(logo_path, width=Inches(1.8))
            p_logo = doc.paragraphs[-1]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_logo.paragraph_format.space_after = Pt(12)
        except Exception:
            pass

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(6)
    
    run_title = p_title.add_run(title.upper())
    run_title.bold = True
    run_title.font.name = font_name
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = PRIMARY_COLOR

    # Subtitle divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(16)
    r_div = p_div.add_run("____________________________________________________")
    r_div.font.color.rgb = SECONDARY_COLOR
    r_div.font.size = Pt(10)

    # Key Terms Summary Table (if provided)
    if summary_table and len(summary_table) > 0:
        p_table_head = doc.add_paragraph()
        p_table_head.paragraph_format.space_before = Pt(10)
        p_table_head.paragraph_format.space_after = Pt(6)
        r_th = p_table_head.add_run("SUMMARY OF KEY TRANSACTION TERMS")
        r_th.bold = True
        r_th.font.name = font_name
        r_th.font.size = Pt(11)
        r_th.font.color.rgb = PRIMARY_COLOR

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Term / Covenant"
        hdr_cells[1].text = "Agreed Specifics"
        hdr_cells[0].width = Inches(2.2)
        hdr_cells[1].width = Inches(4.3)
        
        for c in hdr_cells:
            _set_cell_background(c, "1E293B")
            _set_cell_margins(c, top=140, bottom=140, left=180, right=180)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.name = font_name
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        row_idx = 0
        for k, v in summary_table.items():
            row_cells = table.add_row().cells
            row_cells[0].width = Inches(2.2)
            row_cells[1].width = Inches(4.3)
            row_cells[0].text = str(k)
            row_cells[1].text = str(v)
            
            fill = LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
            for i, c in enumerate(row_cells):
                _set_cell_background(c, fill)
                _set_cell_margins(c, top=100, bottom=100, left=150, right=150)
                _set_cell_borders(c, BORDER_HEX)
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.name = font_name
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = PRIMARY_COLOR if i == 0 else BODY_COLOR
                        if i == 0:
                            r.font.bold = True
            row_idx += 1

        doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Document Body Processing
    cleaned_content = content.replace("```markdown", "").replace("```", "").strip()
    raw_lines = cleaned_content.split("\n")

    in_signature_section = False
    
    for raw_line in raw_lines:
        line = raw_line.strip()
        if not line:
            continue

        # Detect signatures block
        if any(keyword in line.upper() for keyword in ["IN WITNESS WHEREOF", "SIGNATURE PAGE", "EXECUTED AS AN AGREEMENT"]):
            in_signature_section = True
            p_sig = doc.add_paragraph()
            p_sig.paragraph_format.space_before = Pt(20)
            p_sig.paragraph_format.space_after = Pt(8)
            r = p_sig.add_run(line)
            r.bold = True
            r.font.name = font_name
            r.font.size = Pt(11)
            r.font.color.rgb = PRIMARY_COLOR
            continue

        # Major Heading (Markdown # or numbered article e.g. "ARTICLE 1" or "SECTION 1")
        if line.startswith("# ") or (line.isupper() and len(line) < 60 and ("ARTICLE" in line or "SECTION" in line or "AGREEMENT" in line)):
            h_text = line.lstrip("# ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.name = font_name
            r.font.size = Pt(13)
            r.font.color.rgb = PRIMARY_COLOR
            continue

        # Subheading (Markdown ## or ###)
        if line.startswith("## ") or line.startswith("### "):
            sub_text = line.lstrip("# ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sub_text)
            r.bold = True
            r.font.name = font_name
            r.font.size = Pt(11.5)
            r.font.color.rgb = PRIMARY_COLOR
            continue

        # Bullet items
        if line.startswith("* ") or line.startswith("- ") or line.startswith("• "):
            item_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            _format_inline_text(p, item_text, font_name, BODY_COLOR)
            continue

        # Numbered sub-items like 1.1 or (a)
        if re.match(r'^(\d+\.\d+|\([a-z0-9]+\))\s', line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            _format_inline_text(p, line, font_name, BODY_COLOR)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _format_inline_text(p, line, font_name, BODY_COLOR)

    # Optional Clean Dual Signature Block at bottom
    if include_signatures and not in_signature_section:
        _add_signature_table(doc, font_name, PRIMARY_COLOR, SECONDARY_COLOR)

    # Export to memory or file
    target_stream = io.BytesIO()
    doc.save(target_stream)
    docx_bytes = target_stream.getvalue()

    if output_path:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, "wb") as f:
            f.write(docx_bytes)

    return docx_bytes


def _format_inline_text(paragraph, text: str, font_name: str, body_color: RGBColor):
    """Parse inline markdown **bold** or *italic* tags."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        else:
            r = paragraph.add_run(token)
        r.font.name = font_name
        r.font.size = Pt(10.5)
        r.font.color.rgb = body_color


def _add_signature_table(doc, font_name: str, primary_color: RGBColor, slate_color: RGBColor):
    """Add a corporate 2-column signature table."""
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(24)
    p_sig.paragraph_format.space_after = Pt(8)
    p_sig.paragraph_format.keep_with_next = True
    r = p_sig.add_run("IN WITNESS WHEREOF, the Parties hereto have duly executed this Agreement:")
    r.bold = True
    r.font.name = font_name
    r.font.size = Pt(11)
    r.font.color.rgb = primary_color

    sig_table = doc.add_table(rows=5, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    labels = ["PARTY 1 / DISCLOSING PARTY", "PARTY 2 / RECEIVING PARTY"]
    for j in range(2):
        sig_table.rows[0].cells[j].text = labels[j]
        sig_table.rows[1].cells[j].text = "\nBy: _______________________________"
        sig_table.rows[2].cells[j].text = "Name: _____________________________"
        sig_table.rows[3].cells[j].text = "Title: ______________________________"
        sig_table.rows[4].cells[j].text = "Date: ______________________________"

    for row in sig_table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(3.2)
            _set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = font_name
                    r.font.size = Pt(10)
                    r.font.color.rgb = slate_color
